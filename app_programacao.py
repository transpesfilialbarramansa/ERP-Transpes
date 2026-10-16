import io
import base64
import datetime
import hashlib
import json
import psycopg
import pandas as pd
import streamlit as st
import streamlit.components.v1 as components
import re
import openpyxl
import urllib.parse
import smtplib
from email.mime.text import MIMEText
from email.mime.multipart import MIMEMultipart

# Importações para geração de PDF e Excel
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

# ==========================================
# CONFIGURAÇÃO DE E-MAIL E NOTIFICAÇÕES
# ==========================================
EMAIL_DESTINATARIO_GRUPO = "filial.barramansa@transpes.com.br"

def enviar_email_notificacao(assunto, corpo):
    """Envia e-mail automático de notificação do ERP"""
    try:
        msg = MIMEMultipart()
        msg['From'] = f"ERP TRANSPES - Barra Mansa <{st.secrets['EMAIL_REMETENTE']}>"
        msg['To'] = EMAIL_DESTINATARIO_GRUPO
        msg['Subject'] = assunto

        msg.attach(MIMEText(corpo, 'plain', 'utf-8'))

        server = smtplib.SMTP(st.secrets['SMTP_SERVER'], int(st.secrets['SMTP_PORT']))
        server.starttls()
        server.login(st.secrets['EMAIL_REMETENTE'], st.secrets['SENHA_REMETENTE'])
        server.send_message(msg)
        server.close()
    except Exception as e:
        st.error(f"Erro ao enviar e-mail automático: {e}")

def formatar_cpf(valor: str) -> str:
    if not valor:
        return ""
    digits = re.sub(r'\D', '', str(valor))
    if len(digits) == 11:
        return f"{digits[:3]}.{digits[3:6]}.{digits[6:9]}-{digits[9:]}"
    return str(valor).strip()

def formatar_telefone(valor: str) -> str:
    if not valor:
        return ""
    digits = re.sub(r'\D', '', str(valor))
    if len(digits) == 11:
        return f"({digits[:2]}) {digits[2:7]}-{digits[7:]}"
    elif len(digits) == 10:
        return f"({digits[:2]}) {digits[2:6]}-{digits[6:]}"
    return str(valor).strip()

# ==========================================
# DEF_SVGS - DEFINIÇÃO DOS ÍCONES SVG PADRONIZADOS
# ==========================================
SVG_ICONS = {
    "visao_geral": '<svg style="vertical-align: middle; margin-right: 6px;" width="16" height="16" viewBox="0 0 24 24" fill="none" stroke="#D99B26" stroke-width="2" stroke-linecap="round" stroke-linejoin="round"><line x1="18" y1="20" x2="18" y2="10"></line><line x1="12" y1="20" x2="12" y2="4"></line><line x1="6" y1="20" x2="6" y2="14"></line></svg>',
    "programacao": '<svg style="vertical-align: middle; margin-right: 6px;" width="16" height="16" viewBox="0 0 24 24" fill="none" stroke="#D99B26" stroke-width="2" stroke-linecap="round" stroke-linejoin="round"><rect x="3" y="4" width="18" height="18" rx="2" ry="2"></rect><line x1="16" y1="2" x2="16" y2="6"></line><line x1="8" y1="2" x2="8" y2="6"></line><line x1="3" y1="10" x2="21" y2="10"></line></svg>',
    "comercial": '<svg style="vertical-align: middle; margin-right: 6px;" width="16" height="16" viewBox="0 0 24 24" fill="none" stroke="#D99B26" stroke-width="2" stroke-linecap="round" stroke-linejoin="round"><line x1="12" y1="1" x2="12" y2="23"></line><path d="M17 5H9.5a3.5 3.5 0 0 0 0 7h5a3.5 3.5 0 0 1 0 7H6"></path></svg>',
    "expedicao": '<svg style="vertical-align: middle; margin-right: 6px;" width="16" height="16" viewBox="0 0 24 24" fill="none" stroke="#D99B26" stroke-width="2" stroke-linecap="round" stroke-linejoin="round"><line x1="16.5" y1="9.4" x2="7.5" y2="4.21"></line><path d="M21 16V8a2 2 0 0 0-1-1.73l-7-4a2 2 0 0 0-2 0l-7 4A2 2 0 0 0 3 8v8a2 2 0 0 0 1 1.73l7 4a2 2 0 0 0 2 0l7-4A2 2 0 0 0 21 16z"></path><polyline points="3.27 6.96 12 12.01 20.73 6.96"></polyline><line x1="12" y1="22.08" x2="12" y2="12"></line></svg>',
    "operacional": '<svg style="vertical-align: middle; margin-right: 6px;" width="16" height="16" viewBox="0 0 24 24" fill="none" stroke="#D99B26" stroke-width="2" stroke-linecap="round" stroke-linejoin="round"><circle cx="12" cy="12" r="3"></circle><path d="M19.4 15a1.65 1.65 0 0 0 .33 1.82l.06.06a2 2 0 0 1 0 2.83 2 2 0 0 1-2.83 0l-.06-.06a1.65 1.65 0 0 0-1.82-.33 1.65 1.65 0 0 0-1 1.51V21a2 2 0 0 1-2 2 2 2 0 0 1-2-2v-.09A1.65 1.65 0 0 0 9 19.4a1.65 1.65 0 0 0-1.82.33l-.06.06a2 2 0 0 1-2.83 0 2 2 0 0 1 0-2.83l.06-.06a1.65 1.65 0 0 0 .33-1.82 1.65 1.65 0 0 0-1.51-1H3a2 2 0 0 1-2-2 2 2 0 0 1 2-2h.09A1.65 1.65 0 0 0 4.6 9a1.65 1.65 0 0 0-.33-1.82l-.06-.06a2 2 0 0 1 0-2.83 2 2 0 0 1 2.83 0l.06.06a1.65 1.65 0 0 0 1.82.33H9a1.65 1.65 0 0 0 1-1.51V3a2 2 0 0 1 2-2 2 2 0 0 1 2 2v.09a1.65 1.65 0 0 0 1 1.51 1.65 1.65 0 0 0 1.82-.33l.06-.06a2 2 0 0 1 2.83 0 2 2 0 0 1 0 2.83l-.06.06a1.65 1.65 0 0 0-.33 1.82V9a1.65 1.65 0 0 0 1.51 1H21a2 2 0 0 1 2 2 2 2 0 0 1-2 2h-.09a1.65 1.65 0 0 0-1.51 1z"></path></svg>',
    "administracao": '<svg style="vertical-align: middle; margin-right: 6px;" width="16" height="16" viewBox="0 0 24 24" fill="none" stroke="#D99B26" stroke-width="2" stroke-linecap="round" stroke-linejoin="round"><rect x="2" y="7" width="20" height="14" rx="2" ry="2"></rect><path d="M16 21V5a2 2 0 0 0-2-2h-4a2 2 0 0 0-2 2v16"></path></svg>',
    "fornecedores": '<svg style="vertical-align: middle; margin-right: 6px;" width="16" height="16" viewBox="0 0 24 24" fill="none" stroke="#D99B26" stroke-width="2" stroke-linecap="round" stroke-linejoin="round"><rect x="1" y="3" width="15" height="13"></rect><polygon points="16 8 20 8 23 11 23 16 16 16 16 8"></polygon><circle cx="5.5" cy="18.5" r="2.5"></circle><circle cx="18.5" cy="18.5" r="2.5"></circle></svg>',
    "excluir": '<svg style="vertical-align: middle; margin-right: 6px;" width="16" height="16" viewBox="0 0 24 24" fill="none" stroke="#D99B26" stroke-width="2" stroke-linecap="round" stroke-linejoin="round"><polyline points="3 6 5 6 21 6"></polyline><path d="M19 6v14a2 2 0 0 1-2 2H7a2 2 0 0 1-2-2V6m3 0V4a2 2 0 0 1 2-2h4a2 2 0 0 1 2 2v2"></path></svg>',
    "usuarios": '<svg style="vertical-align: middle; margin-right: 6px;" width="16" height="16" viewBox="0 0 24 24" fill="none" stroke="#D99B26" stroke-width="2" stroke-linecap="round" stroke-linejoin="round"><path d="M17 21v-2a4 4 0 0 0-4-4H5a4 4 0 0 0-4 4v2"></path><circle cx="9" cy="7" r="4"></circle><path d="M23 21v-2a4 4 0 0 0-3-3.87"></path><path d="M16 3.13a4 4 0 0 1 0 7.75"></path></svg>',
    "user": '<svg style="vertical-align: middle; margin-right: 4px;" width="14" height="14" viewBox="0 0 24 24" fill="none" stroke="#CBD5E1" stroke-width="2" stroke-linecap="round" stroke-linejoin="round"><path d="M20 21v-2a4 4 0 0 0-4-4H8a4 4 0 0 0-4 4v2"></path><circle cx="12" cy="7" r="4"></circle></svg>',
    "logout": '<svg style="vertical-align: middle; margin-right: 6px;" width="16" height="16" viewBox="0 0 24 24" fill="none" stroke="#94A3B8" stroke-width="2" stroke-linecap="round" stroke-linejoin="round"><path d="M9 21H5a2 2 0 0 1-2-2V5a2 2 0 0 1 2-2h4"></path><polyline points="16 17 21 12 16 7"></polyline><line x1="21" y1="12" x2="9" y2="12"></line></svg>'
}

# ==========================================
# 1. CONFIGURAÇÃO DA PÁGINA
# ==========================================
st.set_page_config(
    page_title="TRANSPES - Sistema de Gestão",
    page_icon="logo_transpes.png",
    layout="wide"
)

# ==========================================
# APLICAÇÃO DE ESTILO E LAYOUT ULTRACOMPACTO
# ==========================================
st.markdown("""
    <style>
    .stApp { background-color: #F4F6F9; }
    .block-container {
        padding-top: 1rem !important; padding-bottom: 1rem !important;
        padding-left: 1.5rem !important; padding-right: 1.5rem !important;
        max-width: 100% !important;
    }
    [data-testid="stSidebar"] {
        background-color: #0B2136 !important;
        min-width: 230px !important; max-width: 230px !important;
    }
    [data-testid="stSidebar"] * { color: #E2E8F0 !important; }
    [data-testid="stSidebarUserContent"] {
        padding-top: 1.5rem !important; padding-left: 1rem !important; padding-right: 1rem !important;
    }
    [data-testid="stSidebar"] .stButton > button {
        background-color: transparent !important; color: #94A3B8 !important;
        border: none !important; text-align: left !important;
        justify-content: flex-start !important; font-weight: 500 !important;
        font-size: 0.85rem !important; border-radius: 6px !important;
        padding: 8px 12px !important; margin-bottom: 4px !important;
        width: 100% !important; transition: all 0.2s ease-in-out !important;
    }
    [data-testid="stSidebar"] .stButton > button:hover {
        background-color: rgba(255, 255, 255, 0.08) !important; color: #FFFFFF !important;
    }
    [data-testid="stSidebar"] .stButton > button[kind="primary"] {
        background-color: #D99B26 !important; color: #0B2136 !important;
        font-weight: 700 !important; box-shadow: 0 2px 6px rgba(0,0,0,0.2) !important;
    }
    h1 { color: #0F172A !important; font-weight: 700 !important; font-size: 1.25rem !important; margin-bottom: 0.5rem !important; }
    h2, h3 { color: #0F172A !important; font-weight: 600 !important; font-size: 0.95rem !important; margin-top: 0.4rem !important; margin-bottom: 0.3rem !important; }
    div[data-baseweb="input"], div[data-baseweb="select"] { min-height: 32px !important; }
    input, select, textarea { font-size: 0.8rem !important; padding: 4px 8px !important; }
    label { font-size: 0.78rem !important; font-weight: 600 !important; margin-bottom: 2px !important; }
    div[data-testid="metric-container"] {
        background-color: #FFFFFF !important; border-radius: 6px !important;
        padding: 6px 12px !important; border: 1px solid #E2E8F0 !important;
        box-shadow: 0px 1px 3px rgba(0, 0, 0, 0.03) !important;
    }
    div[data-testid="metric-container"] label { font-size: 0.72rem !important; }
    div[data-testid="metric-container"] div[data-testid="stMetricValue"] { font-size: 1.1rem !important; font-weight: 700 !important; }
    div[data-testid="stDataFrame"] { font-size: 0.75rem !important; }
    .stButton > button { font-size: 0.82rem !important; padding: 4px 10px !important; }
    hr { margin: 0.6rem 0 !important; }
    header[data-testid="stHeader"] { background: transparent !important; }
    </style>
""", unsafe_allow_html=True)

# ==========================================
# CONEXÃO E BANCO DE DADOS
# ==========================================

def get_connection():
    db_url = st.secrets["DB_URL"]
    if db_url.startswith("postgres://"):
        db_url = db_url.replace("postgres://", "postgresql://", 1)
    return psycopg.connect(db_url)

def hash_senha(senha):
    return hashlib.sha256(senha.encode()).hexdigest()

FORNECEDORES_INICIAIS = [
    ('99', 'ALEXANDRE SILVANO DE MELO', 'ALEMAO MUNCK', 'CAETITE', 'BA', '28.332.244/0001-90', '77 9925-3952'),
    ('108', 'TERMOSOL CONSTRUTORA E COMERCIO LTDA', 'CONSTRUTORA TERMOSOL', 'CAETITE', 'BA', '06.872.066/0001-58', '77 3454-3659'),
    ('43', 'AGM CONSTRUTORA LTDA', 'AGM CONSTRUTORA', 'CALMON', 'BA', '11.051.592/0001-97', '74 36212416'),
    ('151', 'PEZINHO GUINCHO LOCACAO E TRANSPORTE LTDA - ME', 'PEZINHO GUINCHO LOCACAO E TRANSPORTE', 'FEIRA DE SANTANA', 'BA', '13.332.190/0001-96', '75 36234413')
]

@st.cache_resource
def init_db():
    with get_connection() as conn:
        with conn.cursor() as cursor:
            cursor.execute("""
                CREATE TABLE IF NOT EXISTS usuarios (
                    id SERIAL PRIMARY KEY,
                    usuario TEXT UNIQUE NOT NULL,
                    senha TEXT NOT NULL,
                    nome TEXT NOT NULL,
                    perfil TEXT NOT NULL
                )
            """)

            usuarios_iniciais = [
                ("admin", "Transpes@1966", "Administrador do Sistema", "ADMIN"),
                ("deyves.teixeira", "Transpes@26", "Deyves Teixeira", "ADMIN"),
                ("felipe.cesario", "Transpes@26", "Felipe Cesario", "ADMIN")
            ]

            for usr, pwd, nome, perf in usuarios_iniciais:
                cursor.execute("SELECT COUNT(*) FROM usuarios WHERE usuario = %s", (usr,))
                if cursor.fetchone()[0] == 0:
                    cursor.execute(
                        "INSERT INTO usuarios (usuario, senha, nome, perfil) VALUES (%s, %s, %s, %s)",
                        (usr, hash_senha(pwd), nome, perf)
                    )
                else:
                    cursor.execute(
                        "UPDATE usuarios SET senha = %s WHERE usuario = %s",
                        (hash_senha(pwd), usr)
                    )

            cursor.execute("""
                CREATE TABLE IF NOT EXISTS cargas (
                    id SERIAL PRIMARY KEY,
                    numero_carga TEXT UNIQUE NOT NULL,
                    status TEXT DEFAULT 'PENDENTE_COMERCIAL',
                    data_criacao TIMESTAMP DEFAULT CURRENT_TIMESTAMP
                )
            """)

            novas_colunas_cargas = [
                ("numero_set", "TEXT"),
                ("origens_json", "TEXT"),
                ("destinos_json", "TEXT"),
                ("notas_fiscais_comercial", "TEXT"),
                ("receita_frete", "REAL DEFAULT 0.00"),
                ("receita_pedagio", "REAL DEFAULT 0.00"),
                ("receita_taxa_descarga", "REAL DEFAULT 0.00"),
                ("data_previsao_coleta", "TEXT"),
                ("data_previsao_entrega", "TEXT"),
                ("numero_tp", "TEXT"),
                ("nome_motorista", "TEXT"),
                ("cpf_motorista", "TEXT"),
                ("telefone_motorista", "TEXT"),
                ("tipo_motorista", "TEXT"),
                ("tipo_veiculo", "TEXT"),
                ("quantidade_eixos", "INTEGER"),
                ("placa_cavalo", "TEXT"),
                ("placa_carreta", "TEXT"),
                ("valor_rpa", "REAL DEFAULT 0.00"),
                ("data_coleta", "TEXT"),
                ("previsao_descarga", "TEXT"),
                ("observacoes_programacao", "TEXT"),
                ("tipo_pedagio", "TEXT"),
                ("plataforma", "TEXT"),
                ("vinculo", "TEXT")
            ]

            for col_nome, col_tipo in novas_colunas_cargas:
                cursor.execute(f"ALTER TABLE cargas ADD COLUMN IF NOT EXISTS {col_nome} {col_tipo};")

            cursor.execute("""
                CREATE TABLE IF NOT EXISTS carga_expedicao (
                    carga_id INTEGER PRIMARY KEY REFERENCES cargas(id) ON DELETE CASCADE,
                    numero_cte TEXT,
                    numero_viagem TEXT,
                    numero_mdfe TEXT,
                    numero_contrato TEXT,
                    valor_pedagio_pago REAL DEFAULT 0.00,
                    notas_fiscais_expedicao TEXT,
                    data_saida_filial TEXT NOT NULL,
                    observacoes_expedicao TEXT,
                    data_emissao TIMESTAMP DEFAULT CURRENT_TIMESTAMP
                )
            """)

            cursor.execute("""
                CREATE TABLE IF NOT EXISTS carga_operacional (
                    carga_id INTEGER PRIMARY KEY REFERENCES cargas(id) ON DELETE CASCADE,
                    fornecedor_descarga TEXT NOT NULL,
                    custo_fornecedor_descarga REAL DEFAULT 0.00,
                    data_agendamento_descarga TEXT,
                    forma_pagamento TEXT,
                    prazo_pagamento_dias INTEGER,
                    data_registro TIMESTAMP DEFAULT CURRENT_TIMESTAMP
                )
            """)

            cursor.execute("""
                CREATE TABLE IF NOT EXISTS carga_administracao (
                    carga_id INTEGER PRIMARY KEY REFERENCES cargas(id) ON DELETE CASCADE,
                    valor_adiantamento REAL,
                    valor_saldo REAL,
                    comprovante_entregue INTEGER DEFAULT 0,
                    data_liberacao_saldo TEXT,
                    status_pagamento_saldo TEXT DEFAULT 'PENDENTE'
                )
            """)

            cursor.execute("""
                CREATE TABLE IF NOT EXISTS fornecedores (
                    id SERIAL PRIMARY KEY,
                    nr_contrato TEXT,
                    razao_social TEXT NOT NULL,
                    nome_fantasia TEXT,
                    local_atendimento TEXT,
                    uf TEXT,
                    cpf_cnpj TEXT,
                    contato TEXT,
                    data_cadastro TIMESTAMP DEFAULT CURRENT_TIMESTAMP
                )
            """)

            cursor.execute("SELECT COUNT(*) FROM fornecedores")
            if cursor.fetchone()[0] == 0:
                for row in FORNECEDORES_INICIAIS:
                    cursor.execute("""
                        INSERT INTO fornecedores (nr_contrato, razao_social, nome_fantasia, local_atendimento, uf, cpf_cnpj, contato)
                        VALUES (%s, %s, %s, %s, %s, %s, %s)
                    """, row)

            conn.commit()

init_db()

# ==========================================
# FUNÇÕES AUXILIARES
# ==========================================
def get_base64_image(image_path):
    try:
        with open(image_path, "rb") as img_file:
            return base64.b64encode(img_file.read()).decode()
    except Exception:
        return None

def gerar_numero_carga_novo():
    ano = datetime.datetime.now().year
    prefixo = f"TRP-{ano}-"
    
    with get_connection() as conn:
        with conn.cursor() as cursor:
            cursor.execute("""
                SELECT numero_carga 
                FROM cargas 
                WHERE numero_carga LIKE %s
            """, (f"{prefixo}%",))
            
            cargas = cursor.fetchall()
            
            max_seq = 0
            for row in cargas:
                codigo = row[0]
                partes = codigo.split('-')
                if len(partes) == 3 and partes[2].isdigit():
                    seq = int(partes[2])
                    if seq > max_seq:
                        max_seq = seq
            
            proximo_seq = max_seq + 1
            
    return f"{prefixo}{proximo_seq:03d}"

def processar_json_lista(json_str):
    if not json_str:
        return []
    try:
        data = json.loads(json_str)
        return data if isinstance(data, list) else [str(data)]
    except Exception:
        return [s.strip() for s in str(json_str).split(",") if s.strip()]

def formatar_origens_destinos(json_str, apenas_locais=False):
    itens = processar_json_lista(json_str)
    if not itens:
        return "-"
    res = []
    for item in itens:
        if isinstance(item, dict):
            cli = item.get("cliente", "")
            cid = item.get("cidade", "")
            uf = item.get("estado", "")
            if apenas_locais:
                res.append(f"{cid}/{uf}")
            else:
                res.append(f"{cli} ({cid}/{uf})")
        else:
            res.append(str(item))
    return " | ".join(res)

def limpar_formato_json_lista(valor, separador=", "):
    if not valor or str(valor).strip() in ["-", "None", "", "[]"]:
        return "-"
    if isinstance(valor, str) and (valor.startswith("[") and valor.endswith("]")):
        try:
            dados = json.loads(valor)
            if isinstance(dados, list):
                itens = [str(i).strip() for i in dados if str(i).strip()]
                return separador.join(itens) if itens else "-"
        except Exception:
            pass
    limpo = str(valor).replace("[", "").replace("]", "").replace('"', '').replace("'", "").strip()
    return limpo if limpo else "-"

def gerar_excel_geral(df_dados):
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        df_dados.to_excel(writer, index=False, sheet_name="Detalhamento de Cargas")
        workbook = writer.book
        worksheet = writer.sheets["Detalhamento de Cargas"]
        worksheet.views.sheetView[0].showGridLines = True

        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter

        fill_cabecalho = PatternFill(start_color="0F2A4A", end_color="0F2A4A", fill_type="solid")
        font_cabecalho = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        font_corpo = Font(name="Calibri", size=9)

        align_center = Alignment(horizontal="center", vertical="center")
        align_left = Alignment(horizontal="left", vertical="center")
        align_right = Alignment(horizontal="right", vertical="center")

        borda_fina = Side(border_style="thin", color="D9D9D9")
        borda_caixa = Border(left=borda_fina, right=borda_fina, top=borda_fina, bottom=borda_fina)

        cols_moeda = ["Receita Total (R$)", "RPA", "Pedágio Pago", "Custo Descarga", "Custo Total (R$)", "Margem (R$)"]
        cols_centro = ["Nº TP", "Nº SET", "Data Programada", "Status", "CPF Motorista", "Tipo Motorista", "Placa Cavalo", "Placa Carreta", "CT-e", "Viagem", "MDF-e", "Contrato", "Nota Fiscal", "Tipo de Pedágio", "Plataforma", "Vínculo", "Data Agendamento Descarga", "Data Pagamento Saldo"]

        for col_num, col_name in enumerate(df_dados.columns, start=1):
            cell = worksheet.cell(row=1, column=col_num)
            cell.fill = fill_cabecalho
            cell.font = font_cabecalho
            cell.alignment = align_center

        for row_num in range(2, len(df_dados) + 2):
            for col_num, col_name in enumerate(df_dados.columns, start=1):
                cell = worksheet.cell(row=row_num, column=col_num)
                cell.font = font_corpo
                cell.border = borda_caixa
                if col_name in cols_moeda:
                    cell.alignment = align_right
                    cell.number_format = 'R$ #,##0.00'
                elif col_name == "Margem (%)":
                    cell.alignment = align_right
                    cell.number_format = '0%'
                elif col_name in cols_centro:
                    cell.alignment = align_center
                else:
                    cell.alignment = align_left

        for col in worksheet.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or '')
                if len(val_str) > max_len:
                    max_len = len(val_str)
            worksheet.column_dimensions[col_letter].width = max(max_len + 3, 10)

    output.seek(0)
    return output.getvalue()

def gerar_pdf_geral(df):
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(A4), rightMargin=15, leftMargin=15, topMargin=15, bottomMargin=15)
    elements = []
    styles = getSampleStyleSheet()
    
    style_title = ParagraphStyle('TitleStyle', parent=styles['Heading1'], fontName='Helvetica-Bold', fontSize=11, leading=13, alignment=1, textColor=colors.HexColor("#0f2a4a"), spaceAfter=8)
    style_cell = ParagraphStyle('Cell', parent=styles['Normal'], fontName='Helvetica', fontSize=5.5, leading=6.5, alignment=0)
    style_header = ParagraphStyle('Header', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=6, leading=7, textColor=colors.white, alignment=1)

    elements.append(Paragraph("Relatório Geral de Cargas - Transpes", style_title))

    cols_pdf = ["Nº TP", "Nº SET", "Data Programada", "Status", "Motorista", "Origem", "Destino", "Receita Total (R$)", "Custo Total (R$)", "Margem (R$)", "Margem (%)"]
    df_pdf = df.copy()
    
    for c in ["Receita Total (R$)", "Custo Total (R$)", "Margem (R$)"]:
        if c in df_pdf.columns:
            df_pdf[c] = df_pdf[c].apply(lambda x: f"R$ {x:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".") if isinstance(x, (int, float)) else str(x))
    
    if "Margem (%)" in df_pdf.columns:
        df_pdf["Margem (%)"] = df_pdf["Margem (%)"].apply(lambda x: f"{x:.1f}%" if isinstance(x, (int, float)) else str(x))

    table_data = [[Paragraph(col, style_header) for col in cols_pdf]]

    for _, row in df_pdf.iterrows():
        linha = []
        for col in cols_pdf:
            val = str(row[col]) if pd.notnull(row[col]) and str(row[col]) != "" else "-"
            linha.append(Paragraph(val, style_cell))
        table_data.append(linha)

    col_widths = [45, 55, 55, 60, 95, 150, 150, 65, 65, 65, 45]
    pdf_table = Table(table_data, colWidths=col_widths, repeatRows=1)
    pdf_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#0f2a4a")),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#cccccc")),
        ('TOPPADDING', (0, 0), (-1, -1), 2),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 2),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8f9fa")]),
    ]))

    elements.append(pdf_table)
    doc.build(elements)
    return buffer.getvalue()

# ==========================================
# DIÁLOGOS / POP-UPS
# ==========================================
@st.dialog("Programação Registrada!")
def exibir_popup_programacao(num_tp):
    st.success(f"A carga TP **{num_tp}** foi programada com sucesso!")
    if st.button("OK / Nova Programação", type="primary", use_container_width=True):
        st.session_state["exibir_modal_programacao"] = False
        st.session_state["prog_form_version"] = st.session_state.get("prog_form_version", 0) + 1
        st.rerun()

@st.dialog("Dados Comerciais Registrados!")
def exibir_popup_comercial(num_tp):
    st.success(f"Os dados comerciais da carga TP **{num_tp}** foram salvos com sucesso!")
    if st.button("OK / Próximo Comercial", type="primary", use_container_width=True):
        st.session_state["exibir_modal_comercial"] = False
        st.session_state["com_form_version"] = st.session_state.get("com_form_version", 0) + 1
        st.rerun()

@st.dialog("Carga Expedida!")
def exibir_popup_expedicao(num_tp):
    st.success(f"A carga TP **{num_tp}** foi expedida e está **EM TRÂNSITO**!")
    if st.button("OK / Próxima Carga", type="primary", use_container_width=True):
        st.session_state["exibir_modal_expedicao"] = False
        st.session_state["exp_form_version"] = st.session_state.get("exp_form_version", 0) + 1
        st.rerun()

@st.dialog("Agendamento Confirmado!")
def exibir_popup_operacional(num_tp):
    st.success(f"A carga TP **{num_tp}** teve os dados operacionais salvos!")
    if st.button("OK / Próxima Carga", type="primary", use_container_width=True):
        st.session_state["exibir_modal_operacional"] = False
        st.session_state["op_form_version"] = st.session_state.get("op_form_version", 0) + 1
        st.rerun()

@st.dialog("Processo Finalizado!")
def exibir_popup_administracao(num_tp):
    st.success(f"A carga TP **{num_tp}** teve o acerto concluído!")
    if st.button("OK / Concluir", type="primary", use_container_width=True):
        st.session_state["exibir_modal_administracao"] = False
        st.session_state["adm_form_version"] = st.session_state.get("adm_form_version", 0) + 1
        st.rerun()

@st.dialog("Novo Fornecedor")
def modal_cadastrar_fornecedor():
    with st.form("form_modal_fornecedor"):
        f1, f2 = st.columns(2)
        nr_contrato = f1.text_input("Nr. Contrato").strip()
        cpf_cnpj = f2.text_input("CPF/CNPJ*").strip()

        razao_social = st.text_input("Razão Social*").strip().upper()
        nome_fantasia = st.text_input("Nome Fantasia").strip().upper()

        l1, l2 = st.columns([3, 1])
        local_atendimento = l1.text_input("Local Atendimento*").strip().upper()
        uf = l2.text_input("UF*").strip().upper()

        contato = st.text_input("Contato / Telefone").strip()

        btn_salvar_forn = st.form_submit_button("Salvar Fornecedor", type="primary", use_container_width=True)

    if btn_salvar_forn:
        if not razao_social or not local_atendimento or not uf:
            st.error("Preencha os campos obrigatórios.")
        else:
            try:
                with get_connection() as conn:
                    with conn.cursor() as cursor:
                        cursor.execute("""
                            INSERT INTO fornecedores (nr_contrato, razao_social, nome_fantasia, local_atendimento, uf, cpf_cnpj, contato)
                            VALUES (%s, %s, %s, %s, %s, %s, %s)
                        """, (nr_contrato, razao_social, nome_fantasia, local_atendimento, uf, cpf_cnpj, contato))
                        conn.commit()
                st.success("Fornecedor cadastrado com sucesso!")
                st.rerun()
            except Exception as e:
                st.error(f"Erro ao cadastrar fornecedor: {e}")

# ==========================================
# AUTENTICAÇÃO E SESSÃO
# ==========================================
if "logado" not in st.session_state:
    st.session_state["logado"] = False

if not st.session_state["logado"]:
    bg_base64 = get_base64_image("fundo_transpes.png")
    bg_style = f"background-image: url('data:image/jpeg;base64,{bg_base64}'); background-size: cover; background-position: center; background-repeat: no-repeat; background-attachment: fixed;" if bg_base64 else "background-color: #0E2F56;"

    logo_base64 = get_base64_image("logo_transpes.png")
    logo_html = f'<img src="data:image/png;base64,{logo_base64}" style="max-width: 150px; margin-bottom: 8px;">' if logo_base64 else '<div style="color: #0E2F56; font-size: 22px; font-weight: bold;">ERP Transpes</div>'

    st.markdown(
        f"""
        <style>
        [data-testid="stSidebar"], [data-testid="stHeader"] {{ display: none; }}
        .stApp {{ {bg_style} }}
        [data-testid="stMainBlockContainer"] {{ max-width: 380px !important; padding-top: 4rem !important; margin: auto !important; }}
        .login-card {{ background: rgba(255, 255, 255, 0.95); padding: 20px 18px; border-radius: 12px; box-shadow: 0 6px 24px rgba(0, 0, 0, 0.3); text-align: center; margin-bottom: 15px; }}
        </style>
        """,
        unsafe_allow_html=True
    )

    st.markdown(f'<div class="login-card">{logo_html}<div style="color: #555; font-size: 12px;">Acesse o sistema</div></div>', unsafe_allow_html=True)

    with st.form("form_login"):
        usuario_input = st.text_input("Usuário", placeholder="Digite seu usuário")
        senha_input = st.text_input("Senha", type="password", placeholder="Digite sua senha")
        btn_login = st.form_submit_button("Entrar", use_container_width=True, type="primary")

    if btn_login:
        with get_connection() as conn:
            with conn.cursor() as cursor:
                cursor.execute(
                    "SELECT usuario, nome, perfil FROM usuarios WHERE usuario = %s AND senha = %s",
                    (usuario_input.strip().lower(), hash_senha(senha_input))
                )
                usr = cursor.fetchone()
        
        if usr:
            st.session_state["logado"] = True
            st.session_state["usuario"] = usr[0]
            st.session_state["usuario_nome"] = usr[1]
            st.session_state["usuario_perfil"] = usr[2]
            st.rerun()
        else:
            st.error("Usuário ou senha incorretos.")

    st.stop()

# ==========================================
# NAVEGAÇÃO E SIDEBAR ESTILIZADA
# ==========================================
if "menu" not in st.session_state:
    st.session_state["menu"] = "Visão Geral"

menu_selecionado = st.session_state["menu"]
perfil = st.session_state.get("usuario_perfil", "ADMIN")

opcoes_perfil = ["Visão Geral"]
if perfil == "ADMIN":
    opcoes_perfil.extend(["Programação", "Comercial", "Expedição", "Operacional", "Administração", "Fornecedores", "Excluir Cargas", "Usuários"])
elif perfil == "PROGRAMACAO":
    opcoes_perfil.append("Programação")
elif perfil == "COMERCIAL":
    opcoes_perfil.append("Comercial")
elif perfil == "EXPEDICAO":
    opcoes_perfil.append("Expedição")
elif perfil == "OPERACIONAL":
    opcoes_perfil.extend(["Operacional", "Fornecedores"])
elif perfil == "ADMINISTRATIVO":
    opcoes_perfil.extend(["Administração", "Fornecedores"])

logo_base64_sidebar = get_base64_image("transpes_nova_logo_azul.jpg")

with st.sidebar:
    if logo_base64_sidebar:
        st.markdown(f"""
            <div style="text-align: center; padding-bottom: 12px; margin-bottom: 18px; border-bottom: 1px solid rgba(255,255,255,0.08); background-color: #0B2136; border-radius: 6px;">
                <img src="data:image/jpeg;base64,{logo_base64_sidebar}" style="max-width: 150px; height: auto; display: block; margin: 0 auto 6px auto; filter: brightness(1.05) contrast(1.05);">
                <p style="color: #94A3B8 !important; font-size: 0.65rem !important; font-weight: 600 !important; letter-spacing: 1.5px !important; margin: 0 !important; text-transform: uppercase;">SISTEMA DE GESTÃO</p>
            </div>
        """, unsafe_allow_html=True)
    else:
        st.markdown("""
            <div style="text-align: center; padding-bottom: 12px; margin-bottom: 18px; border-bottom: 1px solid rgba(255,255,255,0.08);">
                <h3 style="color: #FFFFFF !important; margin: 0 0 2px 0 !important; font-weight: 800; font-size: 1.1rem; letter-spacing: 1px;">TRANSPES</h3>
                <p style="color: #94A3B8 !important; font-size: 0.65rem !important; font-weight: 600 !important; letter-spacing: 1.5px !important; margin: 0 !important; text-transform: uppercase;">SISTEMA DE GESTÃO</p>
            </div>
        """, unsafe_allow_html=True)

    menus_icones = {
        "Visão Geral": ("visao_geral", "Visão Geral"),
        "Programação": ("programacao", "Programação"),
        "Comercial": ("comercial", "Comercial"),
        "Expedição": ("expedicao", "Expedição"),
        "Operacional": ("operacional", "Operacional"),
        "Administração": ("administracao", "Administração"),
        "Fornecedores": ("fornecedores", "Fornecedores"),
        "Excluir Cargas": ("excluir", "Excluir Cargas"),
        "Usuários": ("usuarios", "Usuários")
    }

    for item in opcoes_perfil:
        chave_svg, rotulo = menus_icones.get(item, ("visao_geral", item))
        if st.button(
            rotulo, 
            type="primary" if menu_selecionado == item else "secondary", 
            use_container_width=True
        ):
            st.session_state["menu"] = item
            st.rerun()

    st.markdown("""
        <div style="margin-top: 25px; padding-top: 12px; border-top: 1px solid rgba(255,255,255,0.08);"></div>
    """, unsafe_allow_html=True)
    
    st.markdown(
        f"<div style='font-size: 0.78rem; color: #CBD5E1; margin-bottom: 8px;'>{SVG_ICONS['user']} <b>{st.session_state.get('usuario_nome', 'Admin')}</b> <span style='color: #64748B;'>({st.session_state.get('usuario_perfil', 'ADMIN')})</span></div>", 
        unsafe_allow_html=True
    )

    if st.button("Sair", use_container_width=True):
        st.session_state.clear()
        st.rerun()

# ==========================================
# PÁGINAS DO SISTEMA
# ==========================================

# 1. VISÃO GERAL
if menu_selecionado == "Visão Geral":
    st.markdown(f"<h1>{SVG_ICONS['visao_geral']} Visão Geral e Relatórios</h1>", unsafe_allow_html=True)

    query = """
        SELECT 
            c.id AS id_carga,
            COALESCE(c.numero_tp, c.numero_carga) AS "Nº TP",
            c.numero_set AS "Nº SET",
            c.data_coleta AS "Data Programada",
            c.status AS "Status",
            COALESCE(c.nome_motorista, '-') AS "Motorista",
            COALESCE(c.cpf_motorista, '-') AS "CPF Motorista",
            COALESCE(c.tipo_motorista, '-') AS "Tipo Motorista",
            COALESCE(c.placa_cavalo, '-') AS "Placa Cavalo",
            COALESCE(c.placa_carreta, '-') AS "Placa Carreta",
            c.origens_json,
            c.destinos_json,
            COALESCE(c.receita_frete, 0.0) AS receita_frete,
            COALESCE(c.receita_pedagio, 0.0) AS receita_pedagio,
            COALESCE(c.receita_taxa_descarga, 0.0) AS receita_taxa_descarga,
            COALESCE(c.valor_rpa, 0.0) AS "RPA",
            COALESCE(e.valor_pedagio_pago, 0.0) AS "Pedágio Pago",
            COALESCE(o.custo_fornecedor_descarga, 0.0) AS "Custo Descarga",
            COALESCE(o.fornecedor_descarga, '-') AS "Fornecedor Descarga",
            COALESCE(o.data_agendamento_descarga, '-') AS "Data Agendamento Descarga",
            COALESCE(e.numero_cte, '-') AS "CT-e",
            COALESCE(e.numero_viagem, '-') AS "Viagem",
            COALESCE(e.numero_mdfe, '-') AS "MDF-e",
            COALESCE(e.numero_contrato, '-') AS "Contrato",
            CASE 
                WHEN e.notas_fiscais_expedicao IS NOT NULL AND e.notas_fiscais_expedicao != '' AND e.notas_fiscais_expedicao != '[]' 
                THEN e.notas_fiscais_expedicao
                ELSE c.notas_fiscais_comercial
            END AS "Nota Fiscal",
            COALESCE(c.tipo_pedagio, '-') AS "Tipo de Pedágio",
            COALESCE(c.plataforma, '-') AS "Plataforma",
            COALESCE(c.vinculo, '-') AS "Vínculo",
            COALESCE(a.data_liberacao_saldo, '-') AS "Data Pagamento Saldo"
        FROM cargas c
        LEFT JOIN carga_expedicao e ON c.id = e.carga_id
        LEFT JOIN carga_operacional o ON c.id = o.carga_id
        LEFT JOIN carga_administracao a ON c.id = a.carga_id
        ORDER BY c.id ASC
    """
    with get_connection() as conn:
        with conn.cursor() as cursor:
            cursor.execute(query)
            dados = cursor.fetchall()
            colunas = [desc.name for desc in cursor.description]
            df = pd.DataFrame(dados, columns=colunas)

    if df.empty:
        st.info("Nenhuma carga cadastrada no momento.")
    else:
        def formatar_real(valor):
            return f"R$ {valor:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")

        df["Origem"] = df["origens_json"].apply(lambda x: formatar_origens_destinos(x))
        df["Destino"] = df["destinos_json"].apply(lambda x: formatar_origens_destinos(x))
        df["Cliente Origem"] = df["origens_json"].apply(lambda x: " | ".join([i.get("cliente","") for i in processar_json_lista(x) if isinstance(i, dict)]))
        df["Cliente Destino"] = df["destinos_json"].apply(lambda x: " | ".join([i.get("cliente","") for i in processar_json_lista(x) if isinstance(i, dict)]))

        df["Receita Total (R$)"] = df["receita_frete"] + df["receita_pedagio"] + df["receita_taxa_descarga"]
        df["Custo Total (R$)"] = df["RPA"] + df["Pedágio Pago"] + df["Custo Descarga"]
        df["Margem (R$)"] = df["Receita Total (R$)"] - df["Custo Total (R$)"]
        df["Margem (%)"] = df.apply(lambda r: (r["Margem (R$)"] / r["Receita Total (R$)"] * 100) if r["Receita Total (R$)"] > 0 else 0.0, axis=1)

        df["CT-e"] = df["CT-e"].apply(lambda x: limpar_formato_json_lista(x, separador=", "))
        df["Viagem"] = df["Viagem"].apply(lambda x: limpar_formato_json_lista(x, separador=", "))
        df["MDF-e"] = df["MDF-e"].apply(lambda x: limpar_formato_json_lista(x, separador=", "))
        df["Nota Fiscal"] = df["Nota Fiscal"].apply(lambda x: limpar_formato_json_lista(x, separador=" / "))

        cols_final = [
            "Nº TP", "Nº SET", "Data Programada", "Status", "Motorista", 
            "CPF Motorista", "Tipo Motorista", "Placa Cavalo", "Placa Carreta",
            "Origem", "Destino", "Cliente Origem", "Cliente Destino", 
            "Receita Total (R$)", "RPA", "Pedágio Pago", "Custo Descarga", 
            "Fornecedor Descarga", "Data Agendamento Descarga", "CT-e", 
            "Viagem", "MDF-e", "Contrato", "Nota Fiscal", "Tipo de Pedágio", 
            "Plataforma", "Vínculo", "Custo Total (R$)", "Margem (R$)", 
            "Margem (%)", "Data Pagamento Saldo"
        ]

        rec_tot = df["Receita Total (R$)"].sum()
        custo_tot = df["Custo Total (R$)"].sum()
        margem_tot = df["Margem (R$)"].sum()

        m1, m2, m3, m4, m5 = st.columns
